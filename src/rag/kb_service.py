"""知识库上传摄取：文档抽取 / 音视频转写 → 文本 → 入 RAG(Chroma+BM25) + 注册表。

- 文档(txt/md/pdf/docx)：抽取文本(同步, 快)
- 音视频(mp3/wav/m4a/mp4…)：ffmpeg 转 16k 单声道 → 分段 → DashScope ASR 转写(慢, 走后台任务)
- 入库：复用共享的 MeetingKB.store(同一个 Chroma 集合) + 保存文本到 kb_dir(供 BM25) + MD5 去重
- 检索器缓存失效：入库/删除后置 kb._retriever=None, 下次检索自动重建(含新文档)
"""
import glob
import json
import os
import re
import subprocess
import tempfile
import time

from src.utils.config import abs_of, kb_conf
from src.utils.files import md5_of_text, read_text
from src.utils.logger import logger

DOC_EXT = {".txt", ".md", ".pdf", ".docx", ".xlsx"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg"}
VIDEO_EXT = {".mp4", ".mov", ".mkv", ".avi", ".webm"}


def kind_of(name: str) -> str | None:
    ext = os.path.splitext(name)[1].lower()
    if ext in DOC_EXT:
        return "doc"
    if ext in AUDIO_EXT:
        return "audio"
    if ext in VIDEO_EXT:
        return "video"
    return None


def _kb_dir() -> str:
    return abs_of("kb_dir")


def _registry_path() -> str:
    return os.path.join(_kb_dir(), "registry.json")


def _load_registry() -> dict:
    from src.utils import db
    return db.store_load("kb_registry", _registry_path(), {})


def _save_registry(reg: dict) -> None:
    from src.utils import db
    db.store_save("kb_registry", _registry_path(), reg)


def _shared_kb():
    """共享 MeetingKB 单例(与内部知识 Agent 同一实例, 入库后它即可检索)。"""
    from src.agent.tools import _meeting_kb
    return _meeting_kb()


def _safe_stem(name: str) -> str:
    stem = os.path.splitext(os.path.basename(name))[0]
    stem = re.sub(r"[^\w一-鿿.-]", "_", stem)[:40]
    return stem or "doc"


# ---------------- 文本抽取 ----------------

def extract_text(path: str, ext: str) -> str:
    if ext in (".txt", ".md"):
        return read_text(path)
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        return "\n".join((pg.extract_text() or "") for pg in reader.pages)
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        return "\n".join(p.text for p in d.paragraphs)
    if ext == ".xlsx":
        # 公司合同/台账等表格：每个工作表 -> 每行非空单元格用 " | " 拼成一行文本，供 RAG 检索。
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)  # data_only: 取公式计算值
        lines = []
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip() != ""]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                lines.append(f"# {ws.title}")
                lines.extend(rows)
        wb.close()
        return "\n".join(lines)
    return ""


# ---------------- 音视频转写 (ffmpeg + DashScope ASR) ----------------

def _ffmpeg(args: list[str]) -> None:
    subprocess.run(["ffmpeg", "-y", "-hide_banner", "-loglevel", "error", *args],
                   check=True, capture_output=True)


def transcribe(src_path: str) -> str:
    """音视频 → 16k 单声道 wav → 分段 → DashScope ASR → 文本。"""
    from dashscope.audio.asr import Recognition
    model = kb_conf.get("asr_model", "paraformer-realtime-v2")
    seg = int(kb_conf.get("asr_segment_seconds", 60))
    with tempfile.TemporaryDirectory() as td:
        wav = os.path.join(td, "audio.wav")
        _ffmpeg(["-i", src_path, "-vn", "-ar", "16000", "-ac", "1", wav])   # 视频自动抽音轨
        _ffmpeg(["-i", wav, "-f", "segment", "-segment_time", str(seg),
                 "-ar", "16000", "-ac", "1", os.path.join(td, "seg_%04d.wav")])
        segs = sorted(glob.glob(os.path.join(td, "seg_*.wav"))) or [wav]
        texts: list[str] = []
        for s in segs:
            try:
                rec = Recognition(model=model, format="wav", sample_rate=16000, callback=None)
                r = rec.call(s)
                if getattr(r, "status_code", None) != 200:
                    logger.warning(f"[KB] ASR 段失败 status={getattr(r, 'status_code', None)}")
                    continue
                sents = r.get_sentence() or []
                if isinstance(sents, dict):
                    sents = [sents]
                for it in sents:
                    t = it.get("text") if isinstance(it, dict) else None
                    if t:
                        texts.append(t)
            except Exception as e:
                logger.warning(f"[KB] ASR 段异常: {e}")
        return "\n".join(texts).strip()


# ---------------- 入库 / 列表 / 删除 ----------------

def ingest_text(text: str, name: str, kind: str, r2_key: str = "") -> dict:
    from src.utils import db
    text = (text or "").strip()
    if not text:
        raise ValueError("No text content extracted")
    doc_id = md5_of_text(text)
    kb = _shared_kb()
    chunks = kb.store.add_document(text, doc_id=doc_id, metadata={"source": name, "kind": kind})
    text_file = ""
    if db.USE_DB:
        db.kv_set(f"kb_text:{doc_id}", text)        # 文本存 kv(供 BM25 / 查看),不落本地文件
    else:
        os.makedirs(_kb_dir(), exist_ok=True)
        text_file = f"{_safe_stem(name)}__{doc_id[:8]}.txt"
        with open(os.path.join(_kb_dir(), text_file), "w", encoding="utf-8") as f:
            f.write(text)                   # 持久 + 供 BM25(load_chunks 会扫 kb_dir)
    kb._mark(doc_id)                         # MD5 台账(本地, 无害)
    kb._retriever = None                     # 失效检索器缓存
    reg = _load_registry()
    reg[doc_id] = {"id": doc_id, "name": name, "kind": kind, "chars": len(text),
                   "chunks": chunks, "added_at": time.time(), "text_file": text_file, "r2_key": r2_key}
    _save_registry(reg)
    logger.info(f"[KB] 入库 {name} ({kind}) -> {chunks} 分块")
    return reg[doc_id]


def process_upload(raw_path: str, name: str, r2_key: str = "") -> dict:
    """根据类型抽取/转写 → 入库。r2_key: 原始文件在 R2 的 key(可空)。返回注册信息。"""
    ext = os.path.splitext(name)[1].lower()
    kind = kind_of(name)
    if kind == "doc":
        text = extract_text(raw_path, ext)
    elif kind in ("audio", "video"):
        text = transcribe(raw_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
    return ingest_text(text, name, kind, r2_key)


def list_docs() -> list[dict]:
    return sorted(_load_registry().values(), key=lambda d: d.get("added_at", 0), reverse=True)


def get_doc_text(doc_id: str) -> str | None:
    """读取某份资料提取后的文本(AI 实际检索所依据的内容)，供前端"查看"。"""
    from src.utils import db
    info = _load_registry().get(doc_id)
    if not info:
        return None
    if db.USE_DB:
        return db.kv_get(f"kb_text:{doc_id}", None)
    if not info.get("text_file"):
        return None
    path = os.path.join(_kb_dir(), info["text_file"])
    return read_text(path) if os.path.exists(path) else None


def get_original_url(doc_id: str) -> str | None:
    """原始文件的临时下载链接(存在 R2 才有)。"""
    from src.utils import db
    info = _load_registry().get(doc_id)
    if not info or not info.get("r2_key") or not db.USE_R2:
        return None
    return db.r2_presigned_url(info["r2_key"])


def search_chunks(query: str, k: int = 6) -> list[dict]:
    """RAG 检索预览：返回混合检索+重排命中的片段(供前端透明展示"AI看到了什么")。"""
    kb = _shared_kb()
    try:
        docs = kb.retriever().invoke(query)
    except Exception as e:
        logger.warning(f"[KB] 检索预览失败: {e}")
        return []
    out = []
    for d in docs[:k]:
        out.append({
            "source": d.metadata.get("source") or d.metadata.get("doc_id", "") or "内部资料",
            "text": (d.page_content or "").strip(),
        })
    return out


def delete_doc(doc_id: str) -> bool:
    from src.utils import db
    kb = _shared_kb()
    kb.store.delete_document(doc_id)             # 删向量分块
    reg = _load_registry()
    info = reg.pop(doc_id, None)
    if db.USE_DB:
        db.kv_delete(f"kb_text:{doc_id}")        # 删文本
        if info and info.get("r2_key"):
            db.r2_delete(info["r2_key"])         # 删 R2 原件
    elif info and info.get("text_file"):
        tp = os.path.join(_kb_dir(), info["text_file"])
        if os.path.exists(tp):
            os.remove(tp)
    _save_registry(reg)
    # 从 MD5 台账移除(允许日后重新上传)
    try:
        if os.path.exists(kb.ledger_path):
            kept = [ln for ln in read_text(kb.ledger_path).splitlines() if ln.strip() and ln.strip() != doc_id]
            with open(kb.ledger_path, "w", encoding="utf-8") as f:
                f.write("\n".join(kept) + ("\n" if kept else ""))
    except Exception as e:
        logger.warning(f"[KB] 台账清理失败: {e}")
    kb._retriever = None
    logger.info(f"[KB] 已删除文档 {doc_id}")
    return True
